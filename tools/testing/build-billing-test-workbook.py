from __future__ import annotations

import json
import re
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "BILLING_PHASE_1_TO_12_TEST_CHECKLIST.md"
SCREENSHOTS = ROOT / "docs" / "testing" / "screenshots"
OUTPUT = ROOT / "BILLING_PHASE_1_TO_12_TEST_WORKBOOK.xlsx"

FRONTEND_URL = "https://contentagent-social-staging-mkhae4sjlq-as.a.run.app"
ADMIN_URL = "https://contentagent-admin-staging-mkhae4sjlq-as.a.run.app/admin"
API_URL = "https://contentagent-api-staging-mkhae4sjlq-as.a.run.app"

COLORS = {
    "ink": "17201B",
    "paper": "F7F7F2",
    "white": "FFFFFF",
    "green": "147D64",
    "green_light": "DFF3EB",
    "blue": "2166A5",
    "blue_light": "DCEBFA",
    "amber": "C47700",
    "amber_light": "FFF0CC",
    "red": "C62828",
    "red_light": "FCE1E1",
    "grey": "E6E8E6",
    "grey_dark": "66706A",
    "purple": "6F42C1",
}

THIN_GREY = Side(style="thin", color="D8DDD9")
ALL_BORDER = Border(left=THIN_GREY, right=THIN_GREY, top=THIN_GREY, bottom=THIN_GREY)


TUTORIALS = [
    {
        "id": "G01",
        "title": "Sign in and find the billing test areas",
        "phase": "All phases",
        "safety": "FREE",
        "file": "01_frontend_navigation.png",
        "steps": [
            "Open the staging frontend and sign in with a dedicated verified tester account.",
            "Record the visible credit balance before executing any paid case.",
            "Use Research for Apify tests, Create Content for generation tests, and Billing for user usage checks.",
            "Open the separate staging Admin URL for plans, rates, Stripe, ledger, and Finance tests.",
        ],
    },
    {
        "id": "G02",
        "title": "Test Create Content quotes and generation",
        "phase": "Phases 2-4",
        "safety": "FREE until Generate; then PAID API",
        "file": "02_create_content.png",
        "steps": [
            "Open Create Content and select Video, Image, or Text post.",
            "Change the platform, model, duration/dimension, and references; verify the credit quote refreshes.",
            "Use Fill fields with AI only when the AI-fill charge is approved.",
            "Scroll to the final Generate control and compare the displayed credits with Admin -> Billing Rules -> Quote tester.",
            "For paid execution, record the balance and ledger before clicking Generate, then verify reservation, settlement, output, and retained media.",
        ],
    },
    {
        "id": "G02B",
        "title": "Approve the final Create Content charge",
        "phase": "Phases 2-4",
        "safety": "PAID API",
        "file": "02b_create_content_generate.png",
        "steps": [
            "Confirm the selected model, duration/dimension, and every attachment before proceeding.",
            "Capture the displayed credit charge and the starting balance.",
            "Click Generate only for an approved paid test case.",
            "Verify immediate redirection to My Contents and a Generating row before provider work completes.",
            "After completion, compare the final debit, provider task ID, logs, preview, and download.",
        ],
    },
    {
        "id": "G03",
        "title": "Test Research scans and cache behavior",
        "phase": "Phase 5",
        "safety": "PAID API for new scan; FREE for history/cache",
        "file": "03_research.png",
        "steps": [
            "Choose Search by URL for strict profile/page scope or Search by keywords for broad discovery.",
            "Select only supported platform/content-type filters and record the displayed quote.",
            "Run one approved scan; verify Scanning feedback, leave/reopen behavior, results, actor/run IDs, reservation, and settlement.",
            "Open History and repeat the same cached search; verify there is no new Apify run or charge.",
        ],
    },
    {
        "id": "G04",
        "title": "Test Persona crawl and AI mapping",
        "phase": "Phase 5",
        "safety": "PAID API",
        "file": "04_persona.png",
        "steps": [
            "Open Persona and enter public Facebook and/or LinkedIn profile URLs.",
            "Record the composite crawl plus AI-mapping quote, then click Auto-fill for an approved test.",
            "Verify successful source data is combined into the single active Persona and a partial-source failure is visible.",
            "Save, reload, and confirm the profile and media persist without creating duplicate personas.",
        ],
    },
    {
        "id": "G05",
        "title": "Test Automation spending controls",
        "phase": "Phase 6",
        "safety": "FREE for controls; PAID API for Run Now",
        "file": "05_automation.png",
        "steps": [
            "Open Automation and choose Source Watcher or Scheduled Content.",
            "Configure source/platform/content settings, Persona or uploaded references, schedule, and publishing account.",
            "Set maximum credits per run plus daily/monthly limits; verify an insufficient cap blocks provider work.",
            "For an approved paid run, trace crawl, AI, generation, and publish as separate ledger actions.",
            "Trigger the same due job twice and confirm there is no duplicate content, post, or charge.",
        ],
    },
    {
        "id": "G06",
        "title": "Verify the user Billing dashboard",
        "phase": "Phase 7",
        "safety": "FREE; Stripe test for Checkout cases",
        "file": "06_user_billing.png",
        "steps": [
            "Open Billing and record available, reserved, and plan balances.",
            "Filter Usage by module and provider; compare totals and rows with the Admin credit ledger.",
            "Verify active plans/packages only, payment history order, warnings, and mobile layout.",
            "Use Stripe test mode only for subscription/top-up Checkout cases.",
        ],
    },
    {
        "id": "G07",
        "title": "Navigate Admin Billing",
        "phase": "Phases 1, 8-12",
        "safety": "FREE navigation",
        "file": "07_admin_billing_modules.png",
        "steps": [
            "Open the staging Admin URL and select Billing.",
            "Use Billing Plans and Credit Packages for commercial offers.",
            "Use Rate Cards, Provider Mapping, and Billing Rules to verify server-side quotes and enforcement.",
            "Use Stripe Connection and Live Launch only with the mode and approval required by the test case.",
            "Use Credit Ledger and Finance for debit, payment, reconciliation, and close evidence.",
        ],
    },
    {
        "id": "G08",
        "title": "Adjust a test user's credits",
        "phase": "Phases 1 and 8",
        "safety": "FREE",
        "file": "08_admin_user_billing.png",
        "steps": [
            "Open Admin -> Billing -> User Billing and select the dedicated staging user.",
            "Record the current available/reserved/lifetime balance.",
            "Enter a positive or negative integer, an expiry only when required, and a clear test reason.",
            "Save once; verify the exact balance change plus one append-only ledger and audit row.",
            "Reverse the fixture with a new compensating adjustment; never edit an existing ledger row.",
        ],
    },
    {
        "id": "G09",
        "title": "Validate rules with Quote tester",
        "phase": "Phases 1, 2, 8, and 9",
        "safety": "FREE",
        "file": "09_billing_rules.png",
        "steps": [
            "Open Admin -> Billing -> Billing Rules.",
            "Record active minimum balances, limits, concurrency, and refund policy before editing.",
            "In Quote tester select the exact action/model/duration or resolution used by the frontend case.",
            "Click Calculate quote and compare integer credits, rate card, mapping, and variable min/max to the frontend.",
            "Restore changed settings after the test and verify the audit event.",
        ],
    },
    {
        "id": "G10",
        "title": "Review provider pricing changes",
        "phase": "Phase 9",
        "safety": "NETWORK; may contact provider",
        "file": "10_pricing_sync.png",
        "steps": [
            "Open Admin -> Billing -> Pricing Alerts and review the last run before starting another.",
            "Start Kie or Apify sync only with approved network/provider access.",
            "Verify a run record and proposal are created while active rate cards remain unchanged.",
            "Approve or keep manual only after comparing provider evidence; verify quotes update without code deployment.",
        ],
    },
    {
        "id": "G11",
        "title": "Test Stripe connection and catalogue",
        "phase": "Phases 10 and 10.5",
        "safety": "STRIPE TEST",
        "file": "11_stripe_connection.png",
        "steps": [
            "Open Admin -> Billing -> Stripe Connection and keep Payment environment on Test.",
            "Validate the test account and webhook; confirm raw keys and signing secrets are never displayed.",
            "Sync the catalogue and verify three plans plus five packages each have the expected immutable revision/mapping.",
            "Run test Checkout, webhook replay, refund, subscription, and Portal cases with Stripe test cards only.",
        ],
    },
    {
        "id": "G12",
        "title": "Execute controlled live launch checks",
        "phase": "Phase 11",
        "safety": "LIVE MONEY - written approval required",
        "file": "12_live_launch.png",
        "steps": [
            "Do not execute this guide unless written approval, refund owner, and backup are recorded.",
            "Complete readiness evidence, catalogue comparison, preflight, and owner-only rollout in order.",
            "Run only the smallest approved top-up and subscription smoke payments.",
            "Verify exact grants/refunds and approve the redacted report, then return rollout to disabled or the approved audience.",
            "Confirm Emergency stop blocks new Checkout while signed webhooks continue.",
        ],
    },
    {
        "id": "G13",
        "title": "Trace credits in the append-only ledger",
        "phase": "Phases 1, 3, 7, and 8",
        "safety": "FREE",
        "file": "13_credit_ledger.png",
        "steps": [
            "Open Admin -> Billing -> Credit Ledger.",
            "Filter by the exact user, module, and provider from the test action.",
            "Trace related IDs across reservation, release or settlement, activity log, and provider task.",
            "Verify one settlement, no duplicate rows, integer credits, and no secrets in descriptions/metadata.",
        ],
    },
    {
        "id": "G14",
        "title": "Reconcile and close Finance",
        "phase": "Phase 12",
        "safety": "STRIPE READ-ONLY; close requires approval",
        "file": "14_finance.png",
        "steps": [
            "Open Admin -> Billing -> Finance and select the intended test/live environment and date range.",
            "Run Reconcile Stripe only with read access; confirm it never mutates the credit ledger.",
            "Compare KPI totals, provider cost, fees, non-USD settlement, findings, and trace to source records.",
            "Resolve/ignore/reopen only dedicated findings and capture actor/time/note evidence.",
            "Close a period only after review and backup; verify the immutable snapshot/hash and duplicate-close rejection.",
        ],
    },
]


def strip_markdown(value: str) -> str:
    value = value.replace("`", "").replace("**", "")
    return re.sub(r"\s+", " ", value).strip()


def split_expected(text: str) -> tuple[str, str]:
    match = re.search(r"\bExpected:\s*", text, flags=re.IGNORECASE)
    if not match:
        return text.rstrip("."), "Verify the described behavior and capture objective evidence."
    return text[: match.start()].strip().rstrip("."), text[match.end() :].strip()


def infer_priority(section: str, scenario: str, safety: str) -> str:
    combined = f"{section} {scenario} {safety}".lower()
    if any(word in combined for word in ("live money", "secret", "permission", "balance", "ledger", "duplicate", "refund", "webhook", "isolation")):
        return "P0"
    if "layout" in combined or "formatting" in combined:
        return "P2"
    return "P1"


def infer_safety(raw: str, section: str, scenario: str) -> str:
    upper = raw.upper()
    for token in ("LIVE MONEY", "STRIPE READ-ONLY", "STRIPE TEST", "PAID/NETWORK", "PAID API", "FREE/ISOLATED", "ISOLATED", "FREE"):
        if token in upper:
            return token
    combined = f"{section} {scenario}".lower()
    if "npm run" in combined or "mobile" in combined or "health" in combined:
        return "FREE"
    if "live" in combined and ("payment" in combined or "checkout" in combined):
        return "LIVE MONEY"
    return "FREE"


def phase_from(test_id: str, heading: str) -> str:
    match = re.match(r"P(10\.5|\d+)", test_id)
    if match:
        return f"Phase {match.group(1)}"
    if test_id.startswith("SEC"):
        return "Cross-Phase Security"
    if test_id.startswith("REL"):
        return "Cross-Phase Reliability"
    return heading


def tutorial_for(test_id: str, phase: str, scenario: str) -> str:
    if test_id.startswith(("SEC", "REL")):
        return "G01"
    match = re.match(r"P(10\.5|\d+)", test_id)
    number = match.group(1) if match else ""
    mapping = {
        "1": "G08/G09/G13",
        "2": "G02/G02B",
        "3": "G13",
        "4": "G02/G02B",
        "5": "G03/G04",
        "6": "G05",
        "7": "G06/G13",
        "8": "G07/G08/G09/G13",
        "9": "G10",
        "10": "G11",
        "10.5": "G11",
        "11": "G12",
        "12": "G14",
    }
    if number == "5" and "persona" in scenario.lower():
        return "G04"
    return mapping.get(number, "G01")


def preconditions_for(safety: str) -> str:
    common = "Staging revision and dedicated tester identity recorded; starting state captured."
    if safety == "LIVE MONEY":
        return common + " Written approval, live readiness, backup, refund owner, and spending limit recorded."
    if safety == "STRIPE TEST":
        return common + " Stripe Test mode connected with healthy webhook and synced test catalogue."
    if safety == "STRIPE READ-ONLY":
        return common + " Correct Stripe mode selected; reconciliation permission granted; no ledger mutation allowed."
    if safety in ("PAID API", "PAID/NETWORK"):
        return common + " Provider token/quota verified; displayed quote and maximum approved; before-balance captured."
    if "ISOLATED" in safety:
        return common + " Use an isolated local/dedicated fixture and take a backup; never alter a real ledger row."
    return common


def steps_for(scenario: str, safety: str, tutorial: str) -> str:
    paid = " Confirm the displayed cost and approval before the paid action." if safety in ("PAID API", "PAID/NETWORK", "STRIPE TEST", "LIVE MONEY") else ""
    return (
        f"1. Open the area referenced by tutorial {tutorial}. "
        "2. Record the before state and relevant IDs. "
        f"3. Execute: {scenario}.{paid} "
        "4. Capture UI, request/status, ledger/payment/provider/log evidence as applicable. "
        "5. Compare with Expected Result and restore dedicated fixtures when required."
    )


def parse_checklist() -> list[dict[str, str]]:
    current_h2 = "General"
    current_h3 = ""
    counters: defaultdict[str, int] = defaultdict(int)
    tests: list[dict[str, str]] = []

    prefixes = {
        "Test Setup": "SET",
        "Required Test Identities": "IDENT",
        "Required Balance Fixtures": "BAL",
        "Required Provider Fixtures": "PROV",
        "Data Safety Rules": "SAFE",
        "Financial Calculation Fixtures": "CALC",
        "Cross-Phase Security and Isolation": "SEC",
        "Cross-Phase Reliability and Performance": "REL",
        "Mandatory Regression Suite": "REG",
        "Promotion Gate": "GATE",
    }

    for raw_line in SOURCE.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if line.startswith("## "):
            current_h2 = strip_markdown(line[3:])
            current_h3 = ""
            continue
        if line.startswith("### "):
            current_h3 = strip_markdown(line[4:])
            continue
        if not line.startswith("- [ ]"):
            continue

        body = line[5:].strip()
        bold = re.match(r"\*\*(.+?)\*\*:\s*(.*)", body)
        header_tokens = ""
        if bold:
            header_tokens, description = bold.groups()
            parts = header_tokens.split()
            test_id = parts[0]
            priority = next((part for part in parts[1:] if part in ("P0", "P1", "P2")), "")
        else:
            description = strip_markdown(body)
            section_key = current_h3 or current_h2
            prefix = prefixes.get(section_key, prefixes.get(current_h2, "CHK"))
            counters[prefix] += 1
            test_id = f"{prefix}-{counters[prefix]:03d}"
            priority = ""

        description = strip_markdown(description)
        scenario, expected = split_expected(description)
        section = current_h3 or current_h2
        safety = infer_safety(header_tokens, section, scenario)
        priority = priority or infer_priority(section, scenario, safety)
        phase = phase_from(test_id, current_h2)
        tutorial = tutorial_for(test_id, phase, scenario)

        if current_h2 == "Mandatory Regression Suite":
            category = "Automated/Regression"
        elif current_h2 == "Promotion Gate":
            category = "Promotion Gate"
        elif current_h2 == "Financial Calculation Fixtures":
            category = "Financial Invariant"
        elif current_h2.startswith("Cross-Phase"):
            category = "Cross-Phase"
        elif current_h2 == "Test Setup":
            category = "Prerequisite"
        else:
            category = "Acceptance Test"

        tests.append(
            {
                "Test ID": test_id,
                "Phase / Area": phase,
                "Section": section,
                "Category": category,
                "Priority": priority,
                "Cost / Safety": safety,
                "Tutorial": tutorial,
                "Scenario": scenario,
                "Expected Result": expected,
                "Preconditions": preconditions_for(safety),
                "Test Steps": steps_for(scenario, safety, tutorial),
                "Result": "Not run",
                "Actual Result": "",
                "Evidence Link / Path": "",
                "Defect ID": "",
                "Tester": "",
                "Date / Time MYT": "",
                "Environment": "Staging",
                "Commit / Revision": "",
                "Notes": "",
            }
        )
    return tests


def title(ws, text: str, subtitle: str, end_col: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    ws.cell(1, 1, text)
    ws.cell(1, 1).font = Font(size=22, bold=True, color=COLORS["white"])
    ws.cell(1, 1).fill = PatternFill("solid", fgColor=COLORS["ink"])
    ws.cell(1, 1).alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 38
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_col)
    ws.cell(2, 1, subtitle)
    ws.cell(2, 1).font = Font(size=10, color=COLORS["grey_dark"])
    ws.cell(2, 1).fill = PatternFill("solid", fgColor=COLORS["paper"])
    ws.cell(2, 1).alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 28


def style_header(row) -> None:
    for cell in row:
        cell.fill = PatternFill("solid", fgColor=COLORS["green"])
        cell.font = Font(bold=True, color=COLORS["white"])
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = ALL_BORDER


def add_table(ws, start_row: int, end_row: int, end_col: int, name: str) -> None:
    if end_row <= start_row:
        return
    table = Table(displayName=name, ref=f"A{start_row}:{column_letter(end_col)}{end_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True, showColumnStripes=False)
    ws.add_table(table)


def column_letter(value: int) -> str:
    result = ""
    while value:
        value, remainder = divmod(value - 1, 26)
        result = chr(65 + remainder) + result
    return result


def add_start_here(wb: Workbook, tests: list[dict[str, str]]) -> None:
    ws = wb.active
    ws.title = "Start Here"
    title(ws, "ContentAgent Billing QA Workbook", "Tester-facing execution guide for billing Phases 1-12. Generated from the comprehensive master checklist.", 10)
    ws.sheet_view.showGridLines = False

    ws["A4"] = "Environment"
    ws["B4"] = "URL"
    style_header(ws[4][:2])
    urls = [("Staging Frontend", FRONTEND_URL), ("Staging Admin", ADMIN_URL), ("Staging API", API_URL)]
    for row, (label, url) in enumerate(urls, 5):
        ws.cell(row, 1, label)
        ws.cell(row, 2, url)
        ws.cell(row, 2).hyperlink = url
        ws.cell(row, 2).style = "Hyperlink"

    ws["A9"] = "How to execute"
    ws["A9"].font = Font(size=15, bold=True, color=COLORS["ink"])
    instructions = [
        "1. Start with the Guided Tutorials sheet. Red callouts show where each control is located.",
        "2. In Test Cases, filter by Phase, Priority, Cost / Safety, or Tutorial. Complete P0 first.",
        "3. Record before state, execute one case, capture evidence, and select a Result from the dropdown.",
        "4. Create a Defect Log row for every failure and link its Defect ID back to the test case.",
        "5. Never execute PAID API, STRIPE TEST, STRIPE READ-ONLY, or LIVE MONEY cases without the stated approval/setup.",
        "6. Use Sign-Off only after all promotion gates and mandatory regression tests are complete.",
    ]
    for index, instruction in enumerate(instructions, 10):
        ws.merge_cells(start_row=index, start_column=1, end_row=index, end_column=10)
        ws.cell(index, 1, instruction)
        ws.cell(index, 1).alignment = Alignment(wrap_text=True, vertical="top")

    ws["A17"] = "Safety labels"
    ws["A17"].font = Font(size=15, bold=True, color=COLORS["ink"])
    safety = [
        ("FREE", "No provider spend. Local/app validation and safe UI/API checks.", COLORS["green_light"]),
        ("PAID API", "Consumes Kie, Apify, Upload-Post, or another provider balance.", COLORS["amber_light"]),
        ("STRIPE TEST", "Uses Stripe test mode only; no real money.", COLORS["blue_light"]),
        ("STRIPE READ-ONLY", "Reads Stripe records for reconciliation; must not mutate credits.", COLORS["blue_light"]),
        ("LIVE MONEY", "Creates real payments/refunds. Written approval is mandatory.", COLORS["red_light"]),
    ]
    for row, (label, detail, fill) in enumerate(safety, 18):
        ws.cell(row, 1, label)
        ws.cell(row, 1).font = Font(bold=True, color=COLORS["red"] if label == "LIVE MONEY" else COLORS["ink"])
        ws.cell(row, 1).fill = PatternFill("solid", fgColor=fill)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=10)
        ws.cell(row, 2, detail)
        ws.cell(row, 2).fill = PatternFill("solid", fgColor=fill)

    ws["A25"] = "Execution summary"
    ws["A25"].font = Font(size=15, bold=True, color=COLORS["ink"])
    metrics = [
        ("Total checklist rows", len(tests)),
        ("P0 rows", sum(1 for item in tests if item["Priority"] == "P0")),
        ("Paid/provider rows", sum(1 for item in tests if item["Cost / Safety"] in ("PAID API", "PAID/NETWORK"))),
        ("Live-money rows", sum(1 for item in tests if item["Cost / Safety"] == "LIVE MONEY")),
        ("Pass", '=COUNTIF(\'Test Cases\'!L:L,"Pass")'),
        ("Fail", '=COUNTIF(\'Test Cases\'!L:L,"Fail")'),
        ("Blocked", '=COUNTIF(\'Test Cases\'!L:L,"Blocked")'),
        ("Not run", '=COUNTIF(\'Test Cases\'!L:L,"Not run")'),
    ]
    for offset, (label, value) in enumerate(metrics):
        row = 26 + offset // 4 * 3
        col = 1 + (offset % 4) * 2
        ws.cell(row, col, label)
        ws.cell(row, col).font = Font(bold=True, color=COLORS["grey_dark"])
        ws.cell(row + 1, col, value)
        ws.cell(row + 1, col).font = Font(size=18, bold=True, color=COLORS["green"])
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
        ws.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=col + 1)

    ws["A33"] = "Workbook generated"
    ws["B33"] = datetime.now().astimezone().strftime("%Y-%m-%d %H:%M:%S %Z")
    ws["A34"] = "Source checklist"
    ws["B34"] = SOURCE.name
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 34
    for col in range(3, 11):
        ws.column_dimensions[column_letter(col)].width = 15


def add_test_cases(wb: Workbook, tests: list[dict[str, str]]) -> None:
    ws = wb.create_sheet("Test Cases")
    headers = list(tests[0].keys())
    title(ws, "Test Cases", "Filter, execute, record evidence, and link defects. All rows originate from the master checklist.", len(headers))
    header_row = 4
    for col, header in enumerate(headers, 1):
        ws.cell(header_row, col, header)
    style_header(ws[header_row])
    ws.row_dimensions[header_row].height = 34

    for row_index, item in enumerate(tests, header_row + 1):
        for col_index, header in enumerate(headers, 1):
            cell = ws.cell(row_index, col_index, item[header])
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = ALL_BORDER
        ws.row_dimensions[row_index].height = 54

    final_row = header_row + len(tests)
    add_table(ws, header_row, final_row, len(headers), "BillingTestCases")
    ws.freeze_panes = "H5"
    ws.auto_filter.ref = f"A{header_row}:{column_letter(len(headers))}{final_row}"
    ws.sheet_view.showGridLines = False

    widths = [13, 21, 28, 20, 10, 19, 14, 52, 52, 42, 72, 17, 45, 32, 15, 18, 20, 14, 22, 35]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[column_letter(index)].width = width

    result_col = headers.index("Result") + 1
    priority_col = headers.index("Priority") + 1
    safety_col = headers.index("Cost / Safety") + 1
    result_range = f"{column_letter(result_col)}5:{column_letter(result_col)}{final_row}"
    validation = DataValidation(type="list", formula1='"Not run,Pass,Fail,Blocked,Not applicable,Accepted risk"', allow_blank=False)
    validation.error = "Choose a result from the list."
    validation.errorTitle = "Invalid result"
    ws.add_data_validation(validation)
    validation.add(result_range)

    result_colors = {
        "Pass": COLORS["green_light"],
        "Fail": COLORS["red_light"],
        "Blocked": COLORS["amber_light"],
        "Accepted risk": "F4E6FF",
        "Not applicable": COLORS["grey"],
    }
    result_letter = column_letter(result_col)
    for value, color in result_colors.items():
        ws.conditional_formatting.add(
            result_range,
            FormulaRule(formula=[f'${result_letter}5="{value}"'], fill=PatternFill("solid", fgColor=color)),
        )

    priority_letter = column_letter(priority_col)
    ws.conditional_formatting.add(
        f"{priority_letter}5:{priority_letter}{final_row}",
        FormulaRule(formula=[f'${priority_letter}5="P0"'], fill=PatternFill("solid", fgColor=COLORS["red_light"]), font=Font(bold=True, color=COLORS["red"])),
    )
    safety_letter = column_letter(safety_col)
    ws.conditional_formatting.add(
        f"{safety_letter}5:{safety_letter}{final_row}",
        FormulaRule(formula=[f'${safety_letter}5="LIVE MONEY"'], fill=PatternFill("solid", fgColor=COLORS["red_light"]), font=Font(bold=True, color=COLORS["red"])),
    )


def add_tutorials(wb: Workbook) -> None:
    ws = wb.create_sheet("Guided Tutorials")
    title(ws, "Guided Tutorials", "Follow the numbered steps and red screenshot callouts. Screenshots are from staging and did not trigger paid actions.", 13)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 28
    ws.column_dimensions["D"].width = 34
    for col in range(5, 14):
        ws.column_dimensions[column_letter(col)].width = 14

    row = 4
    for guide in TUTORIALS:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=13)
        cell = ws.cell(row, 1, f"{guide['id']}  {guide['title']}")
        cell.fill = PatternFill("solid", fgColor=COLORS["green"])
        cell.font = Font(size=15, bold=True, color=COLORS["white"])
        cell.alignment = Alignment(vertical="center")
        ws.row_dimensions[row].height = 28

        ws.cell(row + 1, 1, "Area")
        ws.cell(row + 1, 2, guide["phase"])
        ws.cell(row + 2, 1, "Safety")
        ws.cell(row + 2, 2, guide["safety"])
        ws.cell(row + 3, 1, "Steps")
        for meta_row in range(row + 1, row + 4):
            ws.cell(meta_row, 1).font = Font(bold=True, color=COLORS["ink"])
            ws.cell(meta_row, 1).fill = PatternFill("solid", fgColor=COLORS["paper"])
            ws.cell(meta_row, 1).border = ALL_BORDER
            ws.cell(meta_row, 2).border = ALL_BORDER
            ws.cell(meta_row, 2).alignment = Alignment(wrap_text=True, vertical="top")
        if "LIVE MONEY" in guide["safety"]:
            ws.cell(row + 2, 2).fill = PatternFill("solid", fgColor=COLORS["red_light"])
            ws.cell(row + 2, 2).font = Font(bold=True, color=COLORS["red"])
        elif "PAID" in guide["safety"] or "NETWORK" in guide["safety"]:
            ws.cell(row + 2, 2).fill = PatternFill("solid", fgColor=COLORS["amber_light"])

        step_start = row + 4
        for offset, step in enumerate(guide["steps"]):
            current = step_start + offset
            ws.merge_cells(start_row=current, start_column=1, end_row=current, end_column=4)
            ws.cell(current, 1, f"{offset + 1}. {step}")
            ws.cell(current, 1).alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(current, 1).border = ALL_BORDER
            ws.row_dimensions[current].height = 34

        image_path = SCREENSHOTS / guide["file"]
        if image_path.exists():
            picture = Image(str(image_path))
            scale = min(820 / picture.width, 515 / picture.height)
            picture.width = int(picture.width * scale)
            picture.height = int(picture.height * scale)
            ws.add_image(picture, f"E{row + 1}")
        else:
            ws.cell(row + 1, 5, f"Missing screenshot: {guide['file']}")
            ws.cell(row + 1, 5).fill = PatternFill("solid", fgColor=COLORS["red_light"])

        block_rows = 27
        for image_row in range(row + 1, row + block_rows):
            ws.row_dimensions[image_row].height = max(ws.row_dimensions[image_row].height or 15, 22)
        row += block_rows + 1
    ws.freeze_panes = "A4"


def add_test_data(wb: Workbook) -> None:
    ws = wb.create_sheet("Test Data")
    title(ws, "Test Data and Preconditions", "Create or record these staging-only fixtures before execution. Never put passwords, raw keys, or card data in this workbook.", 8)
    sections = {
        "Test identities": [
            "Owner/super-admin with all billing, Stripe, and Finance permissions",
            "Admin with billing management but no live-launch approval",
            "Finance admin with reconcile/export/resolve but no close permission",
            "Support/read-only admin",
            "Verified user with normal credits",
            "Verified user with zero credits",
            "Verified user one credit below a representative quote",
            "Disabled and unverified users",
            "Two users in different ownership/workspace contexts",
        ],
        "Balance fixtures": [
            "Zero balance",
            "Exactly equal to quote",
            "One credit below quote",
            "Above quote with active reservation",
            "Grants with different expiry dates",
            "Fully consumed grant",
            "Partially and fully refunded purchase",
            "Admin-added and admin-removed credits",
        ],
        "Provider fixtures": [
            "Successful immediate response",
            "Successful asynchronous task",
            "Reject before acceptance",
            "Accept then fail",
            "Timeout and malformed response",
            "Duplicate callback/poll result",
            "Expired temporary media URL",
            "Apify zero, one, and many results",
            "Stripe test USD and non-USD settlement",
        ],
    }
    headers = ["Category", "Fixture", "Ready?", "Identifier / Reference", "Owner", "Created MYT", "Cleanup Required", "Notes"]
    for col, header in enumerate(headers, 1):
        ws.cell(4, col, header)
    style_header(ws[4])
    row = 5
    for category, fixtures in sections.items():
        for fixture in fixtures:
            values = [category, fixture, "No", "", "", "", "Yes", ""]
            for col, value in enumerate(values, 1):
                ws.cell(row, col, value)
                ws.cell(row, col).border = ALL_BORDER
                ws.cell(row, col).alignment = Alignment(wrap_text=True, vertical="top")
            row += 1
    add_table(ws, 4, row - 1, len(headers), "BillingTestData")
    ready = DataValidation(type="list", formula1='"No,Yes,Not required"')
    cleanup = DataValidation(type="list", formula1='"Yes,No,Completed"')
    ws.add_data_validation(ready)
    ws.add_data_validation(cleanup)
    ready.add(f"C5:C{row - 1}")
    cleanup.add(f"G5:G{row - 1}")
    widths = [22, 48, 15, 32, 22, 22, 20, 40]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[column_letter(col)].width = width
    ws.freeze_panes = "A5"
    ws.sheet_view.showGridLines = False


def add_calculations(wb: Workbook) -> None:
    ws = wb.create_sheet("Calculation Fixtures")
    title(ws, "Financial Calculation Fixtures", "Compare quote, reservation, settlement, and report values. Update Expected only when the active Admin rate card is intentionally changed.", 9)
    headers = ["Fixture", "Provider Cost USD", "Markup", "Expected User Charge USD", "Credits / USD", "Expected Credits", "Actual Credits", "Result", "Evidence / Notes"]
    rows = [
        ("VEO Fast 720p", 0.30, 2, 0.60, 1000, 600),
        ("Seedance 2.0 - 1 second", 0.205, 2, 0.41, 1000, 410),
        ("Seedance 2.0 - 8 seconds", 1.64, 2, 3.28, 1000, 3280),
        ("GPT Image 2 - 1K", 0.03, 2, 0.06, 1000, 60),
        ("GPT Image 2 - 2K", 0.05, 2, 0.10, 1000, 100),
        ("GPT Image 2 - 4K", 0.08, 2, 0.16, 1000, 160),
        ("Nano Banana Pro - 1/2K", 0.09, 2, 0.18, 1000, 180),
        ("Nano Banana Pro - 4K", 0.12, 2, 0.24, 1000, 240),
        ("Fractional rounding", 0.0001, 2, 0.0002, 1000, 1),
    ]
    for col, header in enumerate(headers, 1):
        ws.cell(4, col, header)
    style_header(ws[4])
    for row_index, values in enumerate(rows, 5):
        for col_index, value in enumerate(values, 1):
            ws.cell(row_index, col_index, value)
            ws.cell(row_index, col_index).border = ALL_BORDER
        ws.cell(row_index, 8, f'=IF(G{row_index}="","Not run",IF(G{row_index}=F{row_index},"Pass","Fail"))')
    add_table(ws, 4, 4 + len(rows), len(headers), "CalculationFixtures")
    for col, width in enumerate([34, 20, 12, 25, 18, 20, 18, 15, 45], 1):
        ws.column_dimensions[column_letter(col)].width = width
    for row in range(5, 5 + len(rows)):
        for col in (2, 4):
            ws.cell(row, col).number_format = '$0.0000'
    ws.freeze_panes = "A5"
    ws.sheet_view.showGridLines = False


def add_automated_tests(wb: Workbook) -> None:
    ws = wb.create_sheet("Automated Tests")
    title(ws, "Mandatory Automated and Browser Regression", "Run from the repository root against the intended revision. Record sanitized output and evidence.", 10)
    headers = ["ID", "Type", "Command / Check", "Environment", "Expected", "Result", "Executed MYT", "Tester", "Evidence / Output", "Defect ID"]
    checks = [
        ("REG-001", "Build", "npm run build", "Local/CI", "Exit 0"),
        ("REG-002", "Automated", "npm run test:stripe-phase10.5", "Local/Staging", "All pass"),
        ("REG-003", "UI", "npm run test:stripe-phase10.5-ui", "Staging", "All pass"),
        ("REG-004", "Automated", "npm run test:stripe-phase11", "Local/Staging", "All pass"),
        ("REG-005", "UI", "npm run test:stripe-phase11-ui", "Staging", "All pass"),
        ("REG-006", "Automated", "npm run test:stripe-phase12", "Local/Staging", "All pass"),
        ("REG-007", "UI", "npm run test:stripe-phase12-ui", "Staging", "All pass"),
        ("REG-008", "Browser", "Desktop smoke: Frontend, Billing, Admin Billing, Finance, Checkout guardrails", "Staging", "No errors"),
        ("REG-009", "Browser", "Mobile smoke at approximately 390 x 844", "Staging", "No page overflow or inaccessible action"),
        ("REG-010", "Persistence", "Verify users, balances, ledger, orders, subscriptions, generated history, and API key status after deploy", "Staging", "All retained"),
    ]
    for col, header in enumerate(headers, 1):
        ws.cell(4, col, header)
    style_header(ws[4])
    for row_index, check in enumerate(checks, 5):
        values = [*check, "Not run", "", "", "", ""]
        for col_index, value in enumerate(values, 1):
            ws.cell(row_index, col_index, value)
            ws.cell(row_index, col_index).border = ALL_BORDER
            ws.cell(row_index, col_index).alignment = Alignment(wrap_text=True, vertical="top")
    add_table(ws, 4, 4 + len(checks), len(headers), "AutomatedRegression")
    validation = DataValidation(type="list", formula1='"Not run,Pass,Fail,Blocked,Not applicable,Accepted risk"')
    ws.add_data_validation(validation)
    validation.add(f"F5:F{4 + len(checks)}")
    for col, width in enumerate([13, 16, 58, 18, 32, 16, 22, 20, 48, 16], 1):
        ws.column_dimensions[column_letter(col)].width = width
    ws.freeze_panes = "A5"
    ws.sheet_view.showGridLines = False


def add_defects(wb: Workbook) -> None:
    ws = wb.create_sheet("Defect Log")
    title(ws, "Defect Log", "Create one row for every failed test. Never include secrets, passwords, card data, or unredacted personal data.", 15)
    headers = ["Defect ID", "Test ID", "Severity", "Title", "Status", "Area", "Steps to Reproduce", "Expected", "Actual", "Evidence", "Owner", "Reported MYT", "Target Fix", "Retest Result", "Notes"]
    for col, header in enumerate(headers, 1):
        ws.cell(4, col, header)
    style_header(ws[4])
    for row in range(5, 105):
        for col in range(1, len(headers) + 1):
            ws.cell(row, col).border = ALL_BORDER
            ws.cell(row, col).alignment = Alignment(wrap_text=True, vertical="top")
    add_table(ws, 4, 104, len(headers), "BillingDefects")
    severity = DataValidation(type="list", formula1='"Critical,High,Medium,Low"')
    status = DataValidation(type="list", formula1='"Open,In progress,Ready for retest,Closed,Accepted risk"')
    retest = DataValidation(type="list", formula1='"Not run,Pass,Fail,Blocked"')
    ws.add_data_validation(severity)
    ws.add_data_validation(status)
    ws.add_data_validation(retest)
    severity.add("C5:C104")
    status.add("E5:E104")
    retest.add("N5:N104")
    widths = [15, 14, 14, 35, 20, 24, 55, 40, 40, 36, 20, 22, 22, 18, 40]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[column_letter(col)].width = width
    ws.freeze_panes = "A5"
    ws.sheet_view.showGridLines = False


def add_signoff(wb: Workbook) -> None:
    ws = wb.create_sheet("Sign-Off")
    title(ws, "Staging Promotion Sign-Off", "Production promotion is blocked until every mandatory gate is Pass or has a documented accepted exception.", 8)
    gates = [
        "Every P0 test passed or has an explicitly accepted exception",
        "Every P1 test passed or has an approved exception",
        "No unresolved Critical/High billing or reconciliation defects",
        "Credit balances reconcile to the append-only ledger",
        "Stripe test account, catalogue, Checkout, webhooks, refunds, and Portal are healthy",
        "Provider quotes, reservations, settlements, releases, and retained outputs are verified",
        "Live rollout remains disabled unless Phase 11 acceptance is approved",
        "Mandatory automated and desktop/mobile regression tests passed",
        "Backup ID, rollback commit, deployment run, and staging revision are recorded",
        "Security, privacy, isolation, and secret-redaction checks passed",
    ]
    headers = ["Gate", "Result", "Evidence / Exception", "Owner", "Reviewed MYT", "Approver", "Approval MYT", "Notes"]
    for col, header in enumerate(headers, 1):
        ws.cell(4, col, header)
    style_header(ws[4])
    for row_index, gate in enumerate(gates, 5):
        values = [gate, "Not run", "", "", "", "", "", ""]
        for col_index, value in enumerate(values, 1):
            ws.cell(row_index, col_index, value)
            ws.cell(row_index, col_index).border = ALL_BORDER
            ws.cell(row_index, col_index).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row_index].height = 42
    add_table(ws, 4, 4 + len(gates), len(headers), "PromotionSignoff")
    validation = DataValidation(type="list", formula1='"Not run,Pass,Fail,Blocked,Accepted risk"')
    ws.add_data_validation(validation)
    validation.add(f"B5:B{4 + len(gates)}")
    for col, width in enumerate([58, 18, 50, 22, 22, 22, 22, 42], 1):
        ws.column_dimensions[column_letter(col)].width = width
    ws.freeze_panes = "A5"
    ws.sheet_view.showGridLines = False


def protect_formula_cells(wb: Workbook) -> None:
    ws = wb["Start Here"]
    for row in (27, 30):
        for col in range(1, 9, 2):
            if isinstance(ws.cell(row, col).value, str) and ws.cell(row, col).value.startswith("="):
                ws.cell(row, col).protection = Protection(locked=True)


def validate_workbook(path: Path, expected_tests: int) -> dict[str, int]:
    workbook = load_workbook(path, data_only=False)
    required = {"Start Here", "Test Cases", "Guided Tutorials", "Test Data", "Calculation Fixtures", "Automated Tests", "Defect Log", "Sign-Off"}
    missing = required.difference(workbook.sheetnames)
    if missing:
        raise RuntimeError(f"Workbook missing sheets: {sorted(missing)}")
    actual_tests = workbook["Test Cases"].max_row - 4
    if actual_tests != expected_tests:
        raise RuntimeError(f"Expected {expected_tests} test rows, found {actual_tests}")
    image_count = len(workbook["Guided Tutorials"]._images)
    expected_images = sum(1 for guide in TUTORIALS if (SCREENSHOTS / guide["file"]).exists())
    if image_count != expected_images:
        raise RuntimeError(f"Expected {expected_images} embedded screenshots, found {image_count}")
    return {"test_rows": actual_tests, "screenshots": image_count, "sheets": len(workbook.sheetnames)}


def main() -> None:
    tests = parse_checklist()
    if not tests:
        raise RuntimeError("No checklist items parsed")
    duplicate_ids = [test_id for test_id, count in Counter(item["Test ID"] for item in tests).items() if count > 1]
    if duplicate_ids:
        raise RuntimeError(f"Duplicate test IDs: {duplicate_ids[:10]}")

    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    add_start_here(wb, tests)
    add_test_cases(wb, tests)
    add_tutorials(wb)
    add_test_data(wb)
    add_calculations(wb)
    add_automated_tests(wb)
    add_defects(wb)
    add_signoff(wb)
    protect_formula_cells(wb)
    wb.save(OUTPUT)
    validation = validate_workbook(OUTPUT, len(tests))
    print(json.dumps({"output": str(OUTPUT), **validation}, indent=2))


if __name__ == "__main__":
    main()
